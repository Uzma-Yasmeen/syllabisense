"""
Build the code deliverables from demo.py and its real captured output:

    submission/SyllabiSense_Demo.ipynb   notebook, split into cells, outputs embedded
    submission/CODE_SUBMISSION.pdf       the code and the output it produced, plus the graphs

    python demo.py > submission/demo_output.txt
    python make_code_submission.py

Nothing here re-runs the models: the outputs are the ones demo.py actually printed.
"""
import json
import os
import re
import subprocess
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf8")

NAME, ROLL, BATCH = "Uzma Yasmeen", "2503A51L34", "20 CSE"
TITLE = "SyllabiSense: AI for Course Coverage and Personalized Learning"
SRC, OUTTXT = "demo.py", "submission/demo_output.txt"
# the system itself, appended after the demonstration so the PDF really is "the code"
APPENDIX = [
    ("app.py", "the Streamlit application"),
    ("report.py", "the downloadable PDF / CSV report"),
    ("diagrams.py", "runs the CNN on the evidence pages"),
    ("analyze.py", "command-line version and the shared feature builder"),
    ("loaders.py", "PDF / PowerPoint / Word readers, one common shape"),
    ("syllabus_topics.py", "parses a syllabus into units and topics"),
    ("units.py", "matches material files to syllabus units"),
    ("depth.py", "evidence -> properly / barely / not taught"),
    ("build_dataset.py", "builds the labelled (topic, page) pairs, split by subject"),
    ("features_tfidf.py", "TF-IDF + truncated SVD -> the 387 features"),
    ("train_ann.py", "trains the evidence ANN"),
    ("train_cnn.py", "trains the diagram CNN (transfer learning)"),
    ("train_crossencoder.py", "fine-tunes the transformer"),
    ("train_depth.py", "trains the 3-class depth model"),
    ("evaluate_gold.py", "scores a model against the human-verified labels"),
    ("compare_fusion.py", "compares five ways of deploying the two evidence models"),
]
IPYNB, PDF = "submission/SyllabiSense_Demo.ipynb", "submission/CODE_SUBMISSION.pdf"
DOCX = "submission/CODE_SUBMISSION.docx"
NAVY, BLUE, GREY = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x2E, 0x74, 0xB5), RGBColor(0x59, 0x59, 0x59)
SECTION = re.compile(r"^# -+ (\d+)\. (.+?)\s*$")


def split_source():
    """demo.py -> [(number, title, code), ...] using its own section comments."""
    lines = open(SRC, encoding="utf8").read().splitlines()
    cells, cur = [], None
    for ln in lines:
        m = SECTION.match(ln)
        if m:
            if cur:
                cells.append(cur)
            cur = [int(m.group(1)), m.group(2), []]
        elif cur:
            cur[2].append(ln)
    if cur:
        cells.append(cur)
    head = "\n".join(lines[:lines.index(next(l for l in lines if SECTION.match(l)))])
    for c in cells:
        while c[2] and not c[2][0].strip():
            c[2].pop(0)
        while c[2] and not c[2][-1].strip():
            c[2].pop()
        c[2] = "\n".join(c[2])
    return head, cells


def split_output():
    """The captured run -> one chunk per printed banner, in order."""
    text = open(OUTTXT, encoding="utf8").read()
    parts = re.split(r"\n(?=={20,}\n(?:STEP \d|SUMMARY))", text)
    return [p.strip("\n") for p in parts if p.strip()]


# ------------------------------------------------------------------ notebook
def build_notebook(head, cells, outputs):
    nb = {"cells": [], "metadata": {
        "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
        "language_info": {"name": "python", "version": "3.11"}},
        "nbformat": 4, "nbformat_minor": 5}

    def md(text):
        nb["cells"].append({"cell_type": "markdown", "metadata": {},
                            "source": text.splitlines(keepends=True)})

    def code(src, out=None, n=None):
        cell = {"cell_type": "code", "execution_count": n, "metadata": {}, "outputs": [],
                "source": src.splitlines(keepends=True)}
        if out:
            cell["outputs"].append({"output_type": "stream", "name": "stdout",
                                    "text": [l + "\n" for l in out.splitlines()]})
        nb["cells"].append(cell)

    md(f"# {TITLE}\n\n**{NAME} · {ROLL} · Batch {BATCH}**  \n"
       f"AI & Deep Learning — code demonstration\n\n"
       f"This notebook loads the trained models and recomputes every number reported in the "
       f"proposal, then runs the whole system end to end on a real syllabus and real lecture "
       f"material. No figure below is typed in by hand.\n")
    # the docstring + imports cell
    code(head.strip(), None, 1)
    for i, (num, title, src) in enumerate(cells):
        md(f"## {num}. {title}\n")
        out = outputs[i - 1] if 0 < i <= len(outputs) else None
        code(src, out, i + 1)
    json.dump(nb, open(IPYNB, "w", encoding="utf8"), indent=1)
    print(f"wrote {IPYNB}  ({len(nb['cells'])} cells)")


# ----------------------------------------------------------------------- pdf
def mono(doc, text, size=7.6, shade_bg=None):
    for ln in text.splitlines() or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.12)
        r = p.add_run(ln if ln.strip() else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(size)
        if shade_bg:
            r.font.color.rgb = shade_bg


def rule(doc, text, color=NAVY, size=12, space_before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    pr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "6")
    el.set(qn("w:color"), "2E74B5")
    el.set(qn("w:space"), "2")
    bd.append(el)
    pr.append(bd)
    return p


def build_pdf(head, cells, outputs):
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = Inches(0.7)
    s.left_margin = s.right_margin = Inches(0.7)

    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("SyllabiSense — Code Submission")
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{NAME} | {ROLL} | Batch {BATCH} | Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), "PAGE")
    fp._p.append(f)

    for text, size, bold, color, after in [
            ("SR UNIVERSITY", 15, True, NAVY, 2),
            ("Department of Computer Science & Artificial Intelligence", 10, False, GREY, 12),
            (TITLE, 16, True, NAVY, 4),
            ("AI & Deep Learning — Code Submission", 11, False, BLUE, 4),
            (f"{NAME}  ·  {ROLL}  ·  Batch {BATCH}", 10, False, GREY, 16)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        rr = p.add_run(text)
        rr.bold = bold
        rr.font.size = Pt(size)
        rr.font.color.rgb = color

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rr = p.add_run("This document contains the demonstration program (demo.py) and the output it "
                   "actually produced when run. The program loads the trained models and "
                   "recomputes every figure quoted in the proposal, then runs the complete system "
                   "on a real syllabus and real lecture material. The same code is supplied as "
                   "SyllabiSense_Demo.ipynb.")
    rr.font.size = Pt(10)

    rule(doc, "Program header and imports", BLUE, 11, 16)
    mono(doc, head.strip())

    # sections flow on rather than each starting a fresh page: a forced break left pages that
    # were two thirds empty
    for i, (num, title, src) in enumerate(cells):
        rule(doc, f"{num}. {title}", space_before=20 if i else 14)
        p = doc.add_paragraph()
        r = p.add_run("Code")
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = BLUE
        mono(doc, src)
        out = outputs[i - 1] if 0 < i <= len(outputs) else None
        if out:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            r = p.add_run("Output")
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = BLUE
            mono(doc, out, 7.2)

    doc.add_page_break()
    rule(doc, "Result graphs produced by the training code")
    for path, cap in [("results_tfidf/history.png", "Evidence ANN — accuracy and loss vs epochs"),
                      ("results_tfidf/confusion_matrix.png", "Evidence ANN — confusion matrix"),
                      ("results_tfidf/roc_curve.png", "Evidence ANN — ROC curve"),
                      ("results_cnn/history.png", "Diagram CNN — accuracy and loss vs epochs"),
                      ("results_cnn/confusion_matrix.png", "Diagram CNN — confusion matrix"),
                      ("results_cnn/examples.png", "Diagram CNN — most and least confident pages"),
                      ("results_crossencoder/history.png", "Cross-encoder — loss per epoch"),
                      ("results_fusion/comparison.png",
                       "Five deployment rules compared on the human-verified set")]:
        if not os.path.exists(path):
            continue
        doc.add_picture(path, width=Inches(5.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(cap)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY

    # ------------------------------------------------- the source code itself
    doc.add_page_break()
    rule(doc, "Appendix — complete source code", space_before=0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Every file below is the system itself, not the demonstration. The dataset "
                  "collection scripts (the NPTEL and MIT OpenCourseWare downloaders) and the "
                  "hand-labelling screens are in the repository but left out here, because they "
                  "were run once to build the dataset and are not part of the working system.")
    r.font.size = Pt(10)

    shipped = 0
    for name, what in APPENDIX:
        if not os.path.exists(name):
            print(f"  missing, skipped: {name}")
            continue
        src = open(name, encoding="utf8").read()
        doc.add_page_break()
        rule(doc, name, BLUE, 12, 0)
        p = doc.add_paragraph()
        r = p.add_run(f"{what}  ·  {len(src.splitlines())} lines")
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY
        mono(doc, src, 7.0)
        shipped += 1
    print(f"appendix: {shipped} source files")

    doc.save(DOCX)
    print(f"wrote {DOCX}")
    ps = (f'$w = New-Object -ComObject Word.Application; $w.Visible = $false; '
          f'$d = $w.Documents.Open("{os.path.abspath(DOCX)}"); '
          f'$d.SaveAs([ref]"{os.path.abspath(PDF)}", [ref]17); $d.Close(); $w.Quit()')
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True,
                       capture_output=True, timeout=300)
        print(f"wrote {PDF}")
    except Exception as e:
        print(f"PDF export failed ({type(e).__name__}) — open the .docx and Save As PDF")


if __name__ == "__main__":
    head, cells = split_source()
    outputs = split_output()
    print(f"{len(cells)} sections in {SRC}, {len(outputs)} output blocks captured")
    build_notebook(head, cells, outputs)
    build_pdf(head, cells, outputs)
