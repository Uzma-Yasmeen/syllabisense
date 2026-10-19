"""
One reader for every course-material format.

Each file becomes a list of blocks with the same shape, so everything downstream (search,
the model, the report) never needs to know what the original file was:

    {"source": "Unit3.pptx", "location": "Slide 11", "heading": "Transformers", "text": "..."}

PowerPoint and Word keep structure PDFs threw away: slide titles, speaker notes (often the real
explanation), and Word heading paths. A PDF only gives page numbers.

    python loaders.py some_folder_or_file ...     # prints what it finds
"""
import glob
import os
import re
import sys

MIN_CHARS = 120          # ignore near-empty blocks
PAGE_CHARS = 1500


def _clean(t):
    return re.sub(r"\s+", " ", t or "").strip()


def read_pdf(path):
    import pypdfium2 as pdfium
    from match import is_garbage, is_navigation, strip_references
    out = []
    doc = pdfium.PdfDocument(path)
    for i in range(len(doc)):
        raw = doc[i].get_textpage().get_text_range()
        if is_navigation(raw) or is_garbage(raw):
            continue
        t = _clean(strip_references(raw))
        if len(t) >= MIN_CHARS:
            out.append({"location": f"Page {i+1}", "heading": "", "text": t[:PAGE_CHARS]})
    doc.close()
    return out


def read_pptx(path):
    """One block per slide: title + shape text + speaker notes (the notes often hold the
    explanation that the bullets leave out)."""
    from pptx import Presentation
    out = []
    pres = Presentation(path)
    for i, slide in enumerate(pres.slides, 1):
        title, body = "", []
        for sh in slide.shapes:
            if not getattr(sh, "has_text_frame", False):
                continue
            txt = _clean(sh.text_frame.text)
            if not txt:
                continue
            if sh == slide.shapes.title and not title:
                title = txt
            else:
                body.append(txt)
            if getattr(sh, "has_table", False):
                body += [_clean(" | ".join(c.text for c in row.cells)) for row in sh.table.rows]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
        text = " ".join(x for x in [title] + body + ([f"Speaker notes: {notes}"] if notes else []) if x)
        if len(text) >= MIN_CHARS:
            out.append({"location": f"Slide {i}", "heading": title, "text": text[:PAGE_CHARS]})
    return out


def read_docx(path):
    """Blocks follow the document's own headings, so evidence can be cited as a section."""
    import docx
    d = docx.Document(path)
    blocks, cur, heading, path_stack = [], [], "", []

    def flush():
        t = _clean(" ".join(cur))
        if len(t) >= MIN_CHARS:
            blocks.append({"location": " > ".join(path_stack) or f"Part {len(blocks)+1}",
                           "heading": heading, "text": t[:PAGE_CHARS]})
        cur.clear()

    for p in d.paragraphs:
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            flush()
            lvl = int(re.sub(r"\D", "", style) or 1)
            path_stack[:] = path_stack[:lvl - 1] + [_clean(p.text)]
            heading = _clean(p.text)
        elif _clean(p.text):
            cur.append(p.text)
            if sum(len(x) for x in cur) > PAGE_CHARS:
                flush()
    flush()
    for tbl in d.tables:
        rows = [" | ".join(c.text for c in r.cells) for r in tbl.rows]
        t = _clean(" ".join(rows))
        if len(t) >= MIN_CHARS:
            blocks.append({"location": "Table", "heading": "", "text": t[:PAGE_CHARS]})
    return blocks


def read_text(path):
    t = open(path, encoding="utf8", errors="ignore").read()
    out, buf = [], []
    for para in t.split("\n\n"):
        buf.append(para)
        if sum(len(x) for x in buf) > PAGE_CHARS:
            out.append({"location": f"Part {len(out)+1}", "heading": "", "text": _clean(" ".join(buf))[:PAGE_CHARS]})
            buf = []
    if buf and len(_clean(" ".join(buf))) >= MIN_CHARS:
        out.append({"location": f"Part {len(out)+1}", "heading": "", "text": _clean(" ".join(buf))[:PAGE_CHARS]})
    return out


READERS = {".pdf": read_pdf, ".pptx": read_pptx, ".ppt": read_pptx,
           ".docx": read_docx, ".txt": read_text, ".md": read_text}
SUPPORTED = sorted(READERS)


def unreadable(paths, frac=0.6):
    """PDFs that are mostly scanned images rather than text.

    load() quietly drops pages with no text layer, which is the right thing to do for the odd
    cover page but the wrong thing to do silently for a whole file: every topic in it would be
    reported as not taught, when really nothing could be read. Returns
    [(path, pages_without_text, total_pages), ...] so the caller can say so.
    """
    import pypdfium2 as pdfium
    out = []
    for p in paths:
        if not p.lower().endswith(".pdf"):
            continue
        try:
            doc = pdfium.PdfDocument(p)
        except Exception:
            continue
        n = len(doc)
        blank = sum(1 for i in range(n)
                    if len(doc[i].get_textpage().get_text_range().strip()) < MIN_CHARS)
        doc.close()
        if n and blank / n >= frac:
            out.append((p, blank, n))
    return out


def load(paths):
    """paths: files and/or folders -> list of blocks with source/location/heading/text."""
    files = []
    for p in paths:
        if os.path.isdir(p):
            for ext in READERS:
                files += glob.glob(os.path.join(p, "**", f"*{ext}"), recursive=True)
        else:
            files.append(p)
    blocks = []
    for f in sorted(set(files)):
        ext = os.path.splitext(f)[1].lower()
        if ext not in READERS:
            print(f"  skipped (unsupported {ext}): {os.path.basename(f)}")
            continue
        try:
            got = READERS[ext](f)
        except Exception as e:
            print(f"  skipped {os.path.basename(f)}: {type(e).__name__} {e}")
            continue
        for b in got:
            b["source"] = os.path.basename(f)   # the name shown to the user
            b["path"] = f                       # kept so the page can be rendered again later
        blocks += got
    return blocks


if __name__ == "__main__":
    import collections
    sys.stdout.reconfigure(encoding="utf-8")
    bl = load(sys.argv[1:])
    print(f"\n{len(bl)} blocks from {len(set(b['source'] for b in bl))} file(s)")
    for src, n in collections.Counter(b["source"] for b in bl).most_common(10):
        print(f"  {n:5d}  {src}")
    for b in bl[:3]:
        print(f"\n[{b['source']} — {b['location']}] {b['heading']}\n  {b['text'][:200]}...")
